from pathlib import Path

import pytest

from app.ingestion.parsers.docx import _heading_level
from app.ingestion.parsers.hierarchy import HierarchyState, normalize_heading_label
from app.ingestion.parsers.markdown import MarkdownParser
from app.ingestion.parsers.txt import TxtParser
from app.ingestion.search_representation import build_search_text
from app.rag.source_metadata import public_source_metadata


def test_hierarchy_state_nests_replaces_and_disambiguates_siblings():
    state = HierarchyState()

    assert state.content_metadata("content") == {"section_role": "content"}
    operations = state.enter_heading(1, "  Operations   ")
    production = state.enter_heading(2, "Production   Access")
    duplicate = state.enter_heading(2, "Servers")
    duplicate_again = state.enter_heading(2, "Servers")
    development = state.enter_heading(1, "Development")

    assert operations["heading_path"] == "Operations"
    assert production["heading_path"] == "Operations > Production Access"
    assert duplicate["heading_path"] == duplicate_again["heading_path"] == "Operations > Servers"
    assert duplicate["structural_parent_id"] != duplicate_again["structural_parent_id"]
    assert development["heading_path"] == "Development"
    assert state.content_metadata("table")["structural_parent_id"] == development["structural_parent_id"]
    with pytest.raises(ValueError):
        state.content_metadata("heading")


def test_hierarchy_ids_are_delimiter_safe_and_normalization_is_stable():
    state = HierarchyState()
    metadata = state.enter_heading(1, " A/B > C#2 ")

    assert normalize_heading_label(" A/B   > C#2 ") == "A/B > C#2"
    assert metadata["structural_parent_id"].startswith("h:1/")
    assert "/A/B" not in metadata["structural_parent_id"]
    assert ">" not in metadata["structural_parent_id"]


def test_docx_heading_styles_and_table_inherit_position(tmp_path: Path):
    from docx import Document

    path = tmp_path / "hierarchy.docx"
    document = Document()
    document.add_paragraph("Preamble")
    document.add_heading("Operations", level=1)
    document.add_paragraph("Production text")
    document.add_heading("Production Access", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Role"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "Engineer"
    document.add_heading("Development", level=1)
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "Name"
    table.cell(1, 0).text = "Bob"
    document.add_paragraph("Looks like a heading", style="Title")
    document.save(path)

    from app.ingestion.parsers.docx import DocxParser

    sections = DocxParser().parse(path)
    headings = [s for s in sections if s.source_metadata.get("section_role") == "heading"]
    assert [s.text for s in headings] == ["Operations", "Production Access", "Development"]
    production_tables = [s for s in sections if s.source_metadata.get("type") == "table" and "Alice" in s.text]
    development_tables = [s for s in sections if s.source_metadata.get("type") == "table" and "Bob" in s.text]
    assert production_tables[0].source_metadata["heading_path"] == "Operations > Production Access"
    assert development_tables[0].source_metadata["heading_path"] == "Development"
    assert all(s.source_metadata["section_role"] == "table" for s in sections if s.source_metadata.get("type", "").startswith("table"))
    preamble = next(s for s in sections if s.text == "Preamble")
    assert "heading_path" not in preamble.source_metadata
    assert _heading_level(type("P", (), {"style": type("S", (), {"name": "Title"})()})()) is None


def test_markdown_hierarchy_ignores_fenced_headings_and_supports_setext(tmp_path: Path):
    path = tmp_path / "runbook.md"
    path.write_text(
        "Preamble\n\n# Operations ##\nProduction text\n\nProduction Access\n---\nrestart service\n\n```md\n# not a heading\n```\n\n## Development\nrestart service\n",
        encoding="utf-8",
    )

    sections = MarkdownParser().parse(path)
    assert [s.text for s in sections] == [
        "Preamble",
        "# Operations ##",
        "Production text",
        "Production Access",
        "restart service\n```md\n# not a heading\n```",
        "## Development",
        "restart service",
    ]
    assert sections[1].source_metadata["heading_path"] == "Operations"
    assert sections[3].source_metadata["heading_path"] == "Operations > Production Access"
    assert sections[4].source_metadata["heading_path"] == "Operations > Production Access"
    assert sections[5].source_metadata["heading_path"] == "Operations > Development"


def test_markdown_fence_closing_requires_whitespace_and_never_forms_setext(tmp_path: Path):
    path = tmp_path / "fences.md"
    path.write_text(
        "```md\n"
        "``` not-a-close\n"
        "# not a heading\n"
        "````python\n"
        "# still not a heading\n"
        "````\n"
        "---\n"
        "# Actual heading\n"
        "outside content\n",
        encoding="utf-8",
    )

    sections = MarkdownParser().parse(path)

    assert [section.text for section in sections] == [
        "```md\n``` not-a-close\n# not a heading\n````python\n# still not a heading\n````\n---",
        "# Actual heading",
        "outside content",
    ]
    assert sections[0].source_metadata["section_role"] == "content"
    assert "heading_path" not in sections[0].source_metadata
    assert sections[1].source_metadata["heading_path"] == "Actual heading"


@pytest.mark.parametrize(
    ("opening", "closing"),
    [
        ("```~python", "```"),
        ("~~~`python", "~~~"),
    ],
)
def test_markdown_mixed_fence_info_string_does_not_change_delimiter_length(
    tmp_path: Path,
    opening: str,
    closing: str,
):
    path = tmp_path / "mixed-fence.md"
    path.write_text(
        f"{opening}\n"
        "# inside code\n"
        f"{closing}\n"
        "# Operations\n"
        "restart service\n",
        encoding="utf-8",
    )

    sections = MarkdownParser().parse(path)

    assert [section.text for section in sections] == [
        f"{opening}\n# inside code\n{closing}",
        "# Operations",
        "restart service",
    ]
    assert "heading_path" not in sections[0].source_metadata
    assert sections[1].source_metadata["heading_path"] == "Operations"


@pytest.mark.parametrize(
    ("marker", "underline"),
    [("#", "---"), ("##", "===")],
)
def test_markdown_empty_atx_marker_cannot_become_setext_heading(
    tmp_path: Path,
    marker: str,
    underline: str,
):
    path = tmp_path / "empty-atx.md"
    path.write_text(f"{marker}\n{underline}\nordinary content\n", encoding="utf-8")

    sections = MarkdownParser().parse(path)

    assert [section.text for section in sections] == [f"{marker}\n{underline}\nordinary content"]
    assert sections[0].source_metadata["section_role"] == "content"
    assert "heading_path" not in sections[0].source_metadata


@pytest.mark.parametrize("marker", ["# #", "## ###"])
def test_markdown_closing_hash_only_atx_marker_stays_content(tmp_path: Path, marker: str):
    path = tmp_path / "closing-hashes.md"
    path.write_text(f"{marker}\nfollowing content\n", encoding="utf-8")

    sections = MarkdownParser().parse(path)

    assert len(sections) == 1
    assert sections[0].text == f"{marker}\nfollowing content"
    assert sections[0].source_metadata["section_role"] == "content"
    assert "heading_path" not in sections[0].source_metadata


def test_markdown_hash_inside_non_empty_atx_label_remains_heading(tmp_path: Path):
    path = tmp_path / "hash-label.md"
    path.write_text("## C#\nfollowing content\n", encoding="utf-8")

    sections = MarkdownParser().parse(path)

    assert sections[0].source_metadata["heading_path"] == "C#"


def test_txt_does_not_infer_hierarchy_and_search_redacts_internal_fields(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("# Plain text\nbody", encoding="utf-8")
    section = TxtParser().parse(path)[0]
    assert len(TxtParser().parse(path)) == 1
    assert "heading_path" not in section.source_metadata
    assert "structural_parent_id" not in section.source_metadata

    search = build_search_text(
        document_name="runbook.md",
        payload_text="restart service",
        source_metadata={
            "heading_path": "Operations > Production",
            "structural_parent_id": "secret",
            "line_start": 9,
        },
    )
    assert "Heading: Operations > Production" in search
    assert "secret" not in search
    assert "line_start" not in search
    public = public_source_metadata({
        "heading_path": "Operations",
        "section_role": "heading",
        "heading_level": 1,
        "structural_parent_id": "secret",
        "line_start": 2,
    })
    assert public == {
        "heading_path": "Operations",
        "section_role": "heading",
        "heading_level": 1,
    }
