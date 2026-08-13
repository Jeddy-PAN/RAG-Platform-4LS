from pathlib import Path
import uuid

import fitz
import pytest
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.ingestion.pipeline import ingest_document_job
from app.ingestion.search_representation import CURRENT_SEARCH_REPRESENTATION_VERSION
from app.models.chunk import Chunk
from app.models.document import (
    Document,
    DocumentSection,
    DocumentStatus,
    IngestionJob,
    IngestionJobStatus,
)
from app.models.project import Project


class FakeEmbeddingProvider:
    def embed_texts(self, texts: list[str]) -> list[list[float]]:
        return [[0.1] * 1024 for _ in texts]


class RecordingEmbeddingProvider:
    def __init__(self) -> None:
        self.inputs: list[list[str]] = []

    def embed_texts(self, texts: list[str]) -> list[list[float]]:
        self.inputs.append(list(texts))
        return [[0.1] * 1024 for _ in texts]


def test_ingestion_pipeline_creates_sections_and_chunks(
    sqlite_session_factory,
    tmp_path: Path,
) -> None:
    """Successful ingestion should index sections and embedded chunks."""

    path = tmp_path / "source.txt"
    path.write_text("alpha beta gamma delta", encoding="utf-8")

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name="Ingestion")
        db.add(project)
        db.flush()
        project_id = project.id
        document = Document(
            project_id=project_id,
            filename="source.txt",
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(job)
        db.commit()

        ingest_document_job(
            db,
            job.id,
            project.id,
            document.id,
            embedding_provider=FakeEmbeddingProvider(),
            chunk_size=3,
            chunk_overlap=1,
        )

        db.refresh(document)
        db.refresh(job)
        sections = db.query(DocumentSection).all()
        chunks = db.query(Chunk).all()

    assert job.status == IngestionJobStatus.completed
    assert document.status == DocumentStatus.indexed
    assert [section.text for section in sections] == ["alpha beta gamma delta"]
    assert [chunk.project_id for chunk in chunks] == [project_id, project_id]
    assert all(chunk.embedding for chunk in chunks)


def test_ingestion_embeds_search_text_and_stores_payload(
    sqlite_session_factory,
    tmp_path: Path,
) -> None:
    """Provider inputs equal search_text while chunk.text stays the payload."""

    path = tmp_path / "source.txt"
    path.write_text("alpha beta gamma", encoding="utf-8")
    provider = RecordingEmbeddingProvider()

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name="Search Text")
        db.add(project)
        db.flush()
        document = Document(
            project_id=project.id,
            filename="source.txt",
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(job)
        db.commit()

        ingest_document_job(
            db,
            job.id,
            project.id,
            document.id,
            embedding_provider=provider,
        )
        db.refresh(document)
        chunks = db.query(Chunk).all()

    assert provider.inputs
    embedded = provider.inputs[0]
    assert len(embedded) == len(chunks)
    for chunk, search_text in zip(chunks, embedded, strict=True):
        assert chunk.search_text == search_text
        assert chunk.text == "alpha beta gamma"
        assert chunk.search_representation_version == CURRENT_SEARCH_REPRESENTATION_VERSION
        assert chunk.search_text.startswith("Document: source.txt")
        assert "alpha beta gamma" in chunk.search_text
    assert document.search_representation_version == CURRENT_SEARCH_REPRESENTATION_VERSION


def test_successful_reindex_leaves_one_active_current_generation(
    sqlite_session_factory,
    tmp_path: Path,
) -> None:
    """A successful reindex keeps a single active generation at the current version."""

    path = tmp_path / "source.txt"
    path.write_text("alpha beta gamma delta", encoding="utf-8")

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name="Reindex Generations")
        db.add(project)
        db.flush()
        document = Document(
            project_id=project.id,
            filename="source.txt",
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        first_job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(first_job)
        db.commit()

        ingest_document_job(
            db,
            first_job.id,
            project.id,
            document.id,
            embedding_provider=FakeEmbeddingProvider(),
        )
        second_job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(second_job)
        db.commit()
        ingest_document_job(
            db,
            second_job.id,
            project.id,
            document.id,
            embedding_provider=FakeEmbeddingProvider(),
        )

        db.refresh(document)
        active = db.query(Chunk).filter(Chunk.is_active.is_(True)).all()
        all_chunks = db.query(Chunk).all()

    assert len(active) == 1
    assert active[0].search_representation_version == CURRENT_SEARCH_REPRESENTATION_VERSION
    assert len(all_chunks) == 2
    assert {
        chunk.search_representation_version
        for chunk in all_chunks
        if chunk.is_active
    } == {CURRENT_SEARCH_REPRESENTATION_VERSION}
    assert document.search_representation_version == CURRENT_SEARCH_REPRESENTATION_VERSION


@pytest.mark.parametrize("suffix", [".md", ".docx"])
def test_hierarchy_formats_persist_heading_metadata_and_search_context(
    sqlite_session_factory,
    tmp_path: Path,
    suffix: str,
) -> None:
    """DOCX and Markdown indexing persists hierarchy context in v3 chunks."""

    path = tmp_path / f"hierarchy{suffix}"
    if suffix == ".md":
        path.write_text("# Operations\nrestart service", encoding="utf-8")
    else:
        from docx import Document as DocxDocument

        doc = DocxDocument()
        doc.add_heading("Operations", level=1)
        doc.add_paragraph("restart service")
        doc.save(path)

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name=f"Hierarchy {suffix}")
        db.add(project)
        db.flush()
        document = Document(
            project_id=project.id,
            filename=path.name,
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(job)
        db.commit()

        ingest_document_job(
            db,
            job.id,
            project.id,
            document.id,
            embedding_provider=FakeEmbeddingProvider(),
        )
        db.refresh(document)
        active = db.query(Chunk).filter(Chunk.document_id == document.id, Chunk.is_active.is_(True)).all()

    assert document.search_representation_version == "canonical-search-v3"
    assert active
    assert any(chunk.source_metadata.get("heading_path") == "Operations" for chunk in active)
    assert any("Heading: Operations" in chunk.search_text for chunk in active)
    assert all(chunk.search_representation_version == CURRENT_SEARCH_REPRESENTATION_VERSION for chunk in active)


def test_pdf_v1_reindex_rebuilds_page_and_raw_table_generation(
    sqlite_session_factory,
    tmp_path: Path,
) -> None:
    path = tmp_path / "legacy.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Operations")
    page.draw_rect(fitz.Rect(72, 110, 308, 200))
    page.draw_line((190, 110), (190, 200))
    page.draw_line((72, 140), (308, 140))
    page.draw_line((72, 170), (308, 170))
    for y, row in [(130, ("Name", "Role")), (160, ("Alice", "Engineer")), (190, ("Bob", "Designer"))]:
        page.insert_text((80, y), row[0])
        page.insert_text((198, y), row[1])
    pdf.save(path)
    pdf.close()

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name="PDF representation migration")
        db.add(project)
        db.flush()
        document = Document(
            project_id=project.id,
            filename=path.name,
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.indexed,
            search_representation_version="canonical-search-v1",
        )
        db.add(document)
        db.flush()
        legacy_chunk = Chunk(
            project_id=project.id,
            document_id=document.id,
            chunk_index=0,
            text="legacy PDF payload",
            token_count=3,
            content_hash="legacy-pdf-generation",
            source_metadata={"format": "pdf"},
            search_text="Document: legacy.pdf\nContent:\nlegacy PDF payload",
            search_representation_version="canonical-search-v1",
            embedding=[0.1] * 1024,
            is_active=True,
        )
        db.add(legacy_chunk)
        job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(job)
        db.commit()

        assert document.needs_reindex is True
        assert job.job_metadata == {}

        ingest_document_job(
            db,
            job.id,
            project.id,
            document.id,
            embedding_provider=FakeEmbeddingProvider(),
        )

        db.refresh(document)
        chunks = db.query(Chunk).filter(Chunk.document_id == document.id).all()
        active = [chunk for chunk in chunks if chunk.is_active]

    assert document.search_representation_version == CURRENT_SEARCH_REPRESENTATION_VERSION
    assert document.needs_reindex is False
    assert legacy_chunk.is_active is False
    assert active
    assert all(chunk.search_representation_version == CURRENT_SEARCH_REPRESENTATION_VERSION for chunk in active)
    assert any("Page: 1" in chunk.search_text for chunk in active)
    table = next(chunk for chunk in active if chunk.source_metadata.get("table_chunk_type") == "table")
    assert table.source_metadata["row_count"] == 3
    assert table.source_metadata["data_row_start"] == 1
    assert table.source_metadata["data_row_end"] == 3


def test_ingestion_pipeline_requires_project_scoped_document(
    sqlite_session_factory,
    tmp_path: Path,
) -> None:
    """Pipeline must not load documents by document_id alone."""

    path = tmp_path / "source.txt"
    path.write_text("alpha", encoding="utf-8")

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name="Correct")
        other_project_id = uuid.uuid4()
        db.add(project)
        db.flush()
        document = Document(
            project_id=project.id,
            filename="source.txt",
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(job)
        db.commit()

        ingest_document_job(
            db,
            job.id,
            other_project_id,
            document.id,
            embedding_provider=FakeEmbeddingProvider(),
        )

        db.refresh(document)
        db.refresh(job)

    assert job.status == IngestionJobStatus.failed
    assert document.status == DocumentStatus.uploaded
    assert "not found" in (job.error_message or "")


@pytest.mark.integration
def test_ingestion_pipeline_writes_chunks_to_postgresql(
    migrated_engine,
    tmp_path: Path,
) -> None:
    """Successful ingestion should insert vector chunks in PostgreSQL."""

    path = tmp_path / "postgres-source.txt"
    path.write_text("postgres alpha beta gamma", encoding="utf-8")

    with Session(migrated_engine) as db:
        project = Project(name="Postgres Ingestion")
        db.add(project)
        db.flush()
        project_id = project.id
        document = Document(
            project_id=project_id,
            filename="postgres-source.txt",
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        job = IngestionJob(project_id=project_id, document_id=document.id)
        db.add(job)
        db.commit()

        ingest_document_job(
            db,
            job.id,
            project_id,
            document.id,
            embedding_provider=FakeEmbeddingProvider(),
        )

        chunk_count = db.query(Chunk).filter(Chunk.project_id == project_id).count()
        search_vector = db.execute(text("SELECT search_vector FROM chunks LIMIT 1")).scalar_one()

    assert chunk_count == 1
    assert search_vector is not None
