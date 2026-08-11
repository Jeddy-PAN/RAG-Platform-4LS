from pathlib import Path
import uuid

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy.orm import Session

from app.ingestion.pipeline import ingest_document_job
from app.models.document import Document, DocumentStatus, IngestionJob
from app.models.project import Project
from app.models.retrieval import RetrievalMode
from app.rag.retrieval.service import run_retrieval


class FakeEmbeddingProvider:
    def embed_texts(self, texts):
        return [[0.1] * 1024 for _ in texts]


def _workbook(path: Path) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Alpha"
    first.append(["Code", "State"])
    first.append(["alpha-01", "active"])
    first_table = Table(displayName="AlphaTable", ref="A1:B2")
    first_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2")
    first.add_table(first_table)

    second = workbook.create_sheet("Beta")
    second.append(["Code", "State"])
    second.append(["beta-01", "active"])
    second_table = Table(displayName="BetaTable", ref="A1:B2")
    second_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2")
    second.add_table(second_table)
    workbook.save(path)


def test_xlsx_table_query_selects_the_named_sheet_table(
    sqlite_session_factory,
    tmp_path: Path,
) -> None:
    path = tmp_path / "tables.xlsx"
    _workbook(path)

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name=f"XLSX retrieval {uuid.uuid4()}")
        db.add(project)
        db.flush()
        document = Document(
            project_id=project.id,
            filename="tables.xlsx",
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(job)
        db.commit()
        ingest_document_job(
            db,
            job.id,
            project.id,
            document.id,
            embedding_provider=FakeEmbeddingProvider(),
        )

        result = run_retrieval(
            db,
            project_id=project.id,
            query="list all rows in AlphaTable",
            mode=RetrievalMode.keyword,
            top_k=4,
        )

    assert result.table_context is not None
    assert result.table_context.table_index == 0
    assert result.table_context.document_name == "tables.xlsx"
    assert any("alpha-01" in candidate.text for candidate in result.results)
    assert all("beta-01" not in candidate.text for candidate in result.results)


def test_xlsx_wide_fallback_table_query_preserves_adjacent_columns(
    sqlite_session_factory,
    tmp_path: Path,
) -> None:
    path = tmp_path / "wide-fallback.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Metric", "", "", "Value"])
    sheet.append(["cpu", "", "", 80])
    sheet.append(["ram", "", "", 70])
    workbook.save(path)

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name=f"XLSX parallel {uuid.uuid4()}")
        db.add(project)
        db.flush()
        document = Document(
            project_id=project.id,
            filename="wide-fallback.xlsx",
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(job)
        db.commit()
        ingest_document_job(
            db,
            job.id,
            project.id,
            document.id,
            embedding_provider=FakeEmbeddingProvider(),
        )
        result = run_retrieval(
            db,
            project_id=project.id,
            query="show all rows in Metric",
            mode=RetrievalMode.keyword,
            top_k=4,
        )

    assert result.table_context is not None
    assert result.table_context.table_index == 0
    assert any("cpu" in candidate.text and "80" in candidate.text for candidate in result.results)


def test_xlsx_headerless_identifier_and_explicit_headers_remain_retrievable(
    sqlite_session_factory,
    tmp_path: Path,
) -> None:
    path = tmp_path / "header-contracts.xlsx"
    workbook = Workbook()
    fallback = workbook.active
    fallback.title = "Fallback"
    fallback.append(["web-01", "active"])
    fallback.append(["web-02", "down"])
    fallback.append(["web-03", "maintenance"])
    fallback.append([])
    fallback.append(["host", "state"])
    fallback.append(["host-01", "active"])
    fallback.append(["host-02", "down"])
    explicit = workbook.create_sheet("Servers")
    explicit.append(["Server", "IP Address"])
    explicit.append(["api-01", "10.0.0.1"])
    explicit.append(["api-02", "10.0.0.2"])
    explicit.add_table(Table(displayName="ServersTable", ref="A1:B3"))
    workbook.save(path)

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name=f"XLSX headers {uuid.uuid4()}")
        db.add(project)
        db.flush()
        document = Document(
            project_id=project.id,
            filename="header-contracts.xlsx",
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(job)
        db.commit()
        ingest_document_job(
            db,
            job.id,
            project.id,
            document.id,
            embedding_provider=FakeEmbeddingProvider(),
        )
        identifier_result = run_retrieval(
            db,
            project_id=project.id,
            query="web-01",
            mode=RetrievalMode.keyword,
            top_k=4,
        )
        header_result = run_retrieval(
            db,
            project_id=project.id,
            query="show all rows in Server",
            mode=RetrievalMode.keyword,
            top_k=4,
        )
        host_result = run_retrieval(
            db,
            project_id=project.id,
            query="show all rows in host",
            mode=RetrievalMode.keyword,
            top_k=4,
        )

    assert any("web-01 | active" in candidate.text for candidate in identifier_result.results)
    assert host_result.table_context is not None
    assert host_result.table_context.table_index == 1
    assert any("host: host-01" in candidate.text for candidate in host_result.results)
    assert header_result.table_context is not None
    assert header_result.table_context.table_index == 2
    assert any("api-01" in candidate.text for candidate in header_result.results)


def test_xlsx_compact_identifier_first_row_is_retrievable(
    sqlite_session_factory,
    tmp_path: Path,
) -> None:
    path = tmp_path / "compact-identifiers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["aaasiems01", "active"])
    sheet.append(["aaasiems02", "down"])
    sheet.append(["aaasiems03", "maintenance"])
    workbook.save(path)

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name=f"XLSX compact IDs {uuid.uuid4()}")
        db.add(project)
        db.flush()
        document = Document(
            project_id=project.id,
            filename="compact-identifiers.xlsx",
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(job)
        db.commit()
        ingest_document_job(
            db,
            job.id,
            project.id,
            document.id,
            embedding_provider=FakeEmbeddingProvider(),
        )
        result = run_retrieval(
            db,
            project_id=project.id,
            query="aaasiems01",
            mode=RetrievalMode.keyword,
            top_k=4,
        )

    assert any("aaasiems01 | active" in candidate.text for candidate in result.results)
