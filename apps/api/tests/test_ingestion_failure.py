from pathlib import Path

import pytest

from sqlalchemy.orm import Session

from app.ingestion.pipeline import ingest_document_job
from app.ingestion.search_representation import CURRENT_SEARCH_REPRESENTATION_VERSION
from app.models.chunk import Chunk
from app.models.document import Document, DocumentStatus, IngestionJob, IngestionJobStatus
from app.models.project import Project


class FailingEmbeddingProvider:
    def embed_texts(self, texts: list[str]) -> list[list[float]]:
        raise RuntimeError("embedding unavailable")


class StaticEmbeddingProvider:
    def embed_texts(self, texts: list[str]) -> list[list[float]]:
        return [[0.1] * 1024 for _ in texts]


def test_ingestion_failure_marks_statuses(sqlite_session_factory, tmp_path: Path) -> None:
    """Failed first ingestion should mark job and document failed."""

    path = tmp_path / "source.txt"
    path.write_text("alpha beta", encoding="utf-8")

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name="Failure")
        db.add(project)
        db.flush()
        document = Document(
            project_id=project.id,
            filename="source.txt",
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(job)
        db.commit()

        ingest_document_job(
            db,
            job.id,
            project.id,
            document.id,
            embedding_provider=FailingEmbeddingProvider(),
        )

        db.refresh(document)
        db.refresh(job)

    assert job.status == IngestionJobStatus.failed
    assert document.status == DocumentStatus.failed
    assert "embedding unavailable" in (job.error_message or "")
    assert document.search_representation_version is None


@pytest.mark.parametrize("suffix", [".txt", ".md", ".docx"])
def test_failed_reindex_preserves_existing_chunks(
    sqlite_session_factory,
    tmp_path: Path,
    suffix: str,
) -> None:
    """Failed reindex should not delete chunks from a previous good index."""

    path = tmp_path / f"source{suffix}"
    if suffix == ".docx":
        from docx import Document as DocxDocument

        doc = DocxDocument()
        doc.add_heading("Operations", level=1)
        doc.add_paragraph("alpha beta gamma")
        doc.save(path)
    elif suffix == ".md":
        path.write_text("# Operations\nalpha beta gamma", encoding="utf-8")
    else:
        path.write_text("alpha beta gamma", encoding="utf-8")

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name="Reindex")
        db.add(project)
        db.flush()
        document = Document(
            project_id=project.id,
            filename=path.name,
            storage_path=str(path),
            file_size_bytes=path.stat().st_size,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        first_job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(first_job)
        db.commit()

        ingest_document_job(
            db,
            first_job.id,
            project.id,
            document.id,
            embedding_provider=StaticEmbeddingProvider(),
        )
        original_chunk_count = db.query(Chunk).count()
        original_active_count = db.query(Chunk).filter(Chunk.is_active.is_(True)).count()

        second_job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(second_job)
        db.commit()
        ingest_document_job(
            db,
            second_job.id,
            project.id,
            document.id,
            embedding_provider=FailingEmbeddingProvider(),
        )

        db.refresh(document)
        db.refresh(second_job)
        final_chunk_count = db.query(Chunk).count()

    assert original_chunk_count > 0
    assert final_chunk_count == original_chunk_count
    assert second_job.status == IngestionJobStatus.failed
    assert document.status == DocumentStatus.indexed
    # The previous good index and its representation version survive the failure.
    assert document.search_representation_version == CURRENT_SEARCH_REPRESENTATION_VERSION
    active = db.query(Chunk).filter(Chunk.is_active.is_(True)).all()
    assert len(active) == original_active_count
    assert all(chunk.search_representation_version == CURRENT_SEARCH_REPRESENTATION_VERSION for chunk in active)
    if suffix != ".txt":
        assert any(chunk.source_metadata.get("heading_path") == "Operations" for chunk in active)
