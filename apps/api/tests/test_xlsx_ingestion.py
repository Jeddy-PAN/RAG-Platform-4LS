from datetime import date
from pathlib import Path
import uuid
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from app.ingestion.chunker import chunk_sections
from app.ingestion.parsers.xlsx import XlsxParser
from app.ingestion.search_representation import build_search_text
from app.ingestion.pipeline import ingest_document_job
from app.models.chunk import Chunk
from app.models.document import Document, DocumentSection, DocumentStatus, IngestionJob, IngestionJobStatus
from app.models.project import Project
from sqlalchemy.orm import Session


def _save_workbook(path: Path):
    from openpyxl import Workbook

    workbook = Workbook()
    return workbook


def test_sparse_xlsx_preserves_visible_sheets_and_empty_columns(tmp_path: Path) -> None:
    workbook = _save_workbook(tmp_path / "sparse.xlsx")
    sheet = workbook.active
    sheet.title = "Visible"
    sheet.append(["Name", "Gap", "Count", "State"])
    sheet.append(["alpha", None, 10, "enabled"])
    second = workbook.create_sheet("Second")
    second.append(["Other", "Value"])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["secret"])
    workbook.save(tmp_path / "sparse.xlsx")

    sections = XlsxParser().parse(tmp_path / "sparse.xlsx")
    row = next(section for section in sections if section.source_metadata["type"] == "table_row")

    assert row.source_metadata["sheet_name"] == "Visible"
    assert row.source_metadata["column_count"] == 4
    assert row.source_metadata["sheet_row"] == 2
    assert row.source_metadata["values"] == ["alpha", "", "10", "enabled"]
    assert all(section.source_metadata["sheet_name"] != "Hidden" for section in sections)
    assert "alpha |  | Count: 10" in row.text


def test_explicit_table_and_separated_fallback_region_are_distinct(tmp_path: Path) -> None:
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = _save_workbook(tmp_path / "regions.xlsx")
    sheet = workbook.active
    sheet.title = "Inventory"
    sheet.append(["Code", "State"])
    sheet.append(["node-01", "active"])
    sheet.append([])
    sheet.append(["Other", "Value"])
    sheet.append(["alpha", "10"])
    table = Table(displayName="InventoryTable", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    workbook.save(tmp_path / "regions.xlsx")

    sections = XlsxParser().parse(tmp_path / "regions.xlsx")
    tables = [section for section in sections if section.source_metadata["type"] == "table"]

    assert len(tables) == 2
    assert [table.source_metadata["table_range"] for table in tables] == ["A1:B2", "A4:B5"]
    assert tables[0].source_metadata["caption"] == "InventoryTable"
    assert len({table.source_metadata["table_index"] for table in tables}) == 2


def test_explicit_headerless_table_keeps_first_row_as_data(tmp_path: Path) -> None:
    from openpyxl.worksheet.table import Table

    workbook = _save_workbook(tmp_path / "headerless-table.xlsx")
    sheet = workbook.active
    sheet.append(["alpha", "10"])
    sheet.append(["beta", "20"])
    table = Table(displayName="HeaderlessTable", ref="A1:B2")
    table.headerRowCount = 0
    sheet.add_table(table)
    workbook.save(tmp_path / "headerless-table.xlsx")

    sections = XlsxParser().parse(tmp_path / "headerless-table.xlsx")
    tables = [section for section in sections if section.source_metadata["type"] == "table"]
    rows = [section for section in sections if section.source_metadata["type"] == "table_row"]

    assert len(tables) == 1
    assert not any(section.source_metadata["type"] == "table_header" for section in sections)
    assert tables[0].source_metadata["headers"] == []
    assert tables[0].source_metadata["header_source"] == "none"
    assert [row.source_metadata["data_row"] for row in rows] == [1, 2]
    assert rows[0].text == "alpha | 10"
    assert "Columns:" not in build_search_text(
        document_name="headerless-table.xlsx",
        payload_text=rows[0].text,
        source_metadata=rows[0].source_metadata,
    )


def test_unstyled_same_row_fallback_blocks_remain_one_rectangular_region(tmp_path: Path) -> None:
    workbook = _save_workbook(tmp_path / "horizontal-regions.xlsx")
    sheet = workbook.active
    sheet["A1"] = "LeftCode"
    sheet["A2"] = "left-01"
    sheet["D1"] = "RightCode"
    sheet["D2"] = "right-01"
    workbook.save(tmp_path / "horizontal-regions.xlsx")

    sections = XlsxParser().parse(tmp_path / "horizontal-regions.xlsx")
    tables = [section for section in sections if section.source_metadata["type"] == "table"]

    assert len(tables) == 1
    assert tables[0].source_metadata["table_range"] == "A1:D2"
    assert tables[0].source_metadata["column_count"] == 4


def test_blank_interior_column_is_not_removed_or_used_as_separator(tmp_path: Path) -> None:
    workbook = _save_workbook(tmp_path / "blank-column.xlsx")
    sheet = workbook.active
    sheet.title = "Sparse"
    sheet.append(["Left", "", "Right"])
    sheet.append(["alpha", "", "omega"])
    workbook.save(tmp_path / "blank-column.xlsx")

    sections = XlsxParser().parse(tmp_path / "blank-column.xlsx")
    row = next(section for section in sections if section.source_metadata["type"] == "table_row")

    assert row.source_metadata["column_count"] == 3
    assert row.source_metadata["values"] == ["alpha", "", "omega"]
    assert row.text == "Left: alpha |  | Right: omega"


def test_multiple_blank_interior_columns_remain_one_rectangular_region(tmp_path: Path) -> None:
    workbook = _save_workbook(tmp_path / "wide-blank-columns.xlsx")
    sheet = workbook.active
    sheet.append(["Metric", "", "", "Value"])
    sheet.append(["cpu", "", "", 80])
    sheet.append(["ram", "", "", 70])
    workbook.save(tmp_path / "wide-blank-columns.xlsx")

    sections = XlsxParser().parse(tmp_path / "wide-blank-columns.xlsx")
    table = next(section for section in sections if section.source_metadata["type"] == "table")
    rows = [section for section in sections if section.source_metadata["type"] == "table_row"]

    assert table.source_metadata["table_range"] == "A1:D3"
    assert table.source_metadata["column_count"] == 4
    assert [row.source_metadata["values"] for row in rows] == [
        ["cpu", "", "", "80"],
        ["ram", "", "", "70"],
    ]


def test_styled_wide_blank_columns_remain_one_rectangular_region(tmp_path: Path) -> None:
    from openpyxl.styles import PatternFill

    workbook = _save_workbook(tmp_path / "styled-wide-blank-columns.xlsx")
    sheet = workbook.active
    sheet.append(["Metric", "", "", "Value"])
    sheet.append(["cpu", "", "", 80])
    sheet.append(["ram", "", "", 70])
    for row in range(1, 4):
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor="D9EAF7")
        sheet.cell(row, 4).fill = PatternFill("solid", fgColor="FCE4D6")
    workbook.save(tmp_path / "styled-wide-blank-columns.xlsx")

    sections = XlsxParser().parse(tmp_path / "styled-wide-blank-columns.xlsx")
    tables = [section for section in sections if section.source_metadata["type"] == "table"]

    assert len(tables) == 1
    assert tables[0].source_metadata["table_range"] == "A1:D3"
    assert tables[0].source_metadata["column_count"] == 4


def test_merged_value_fills_covered_position_with_provenance(tmp_path: Path) -> None:
    workbook = _save_workbook(tmp_path / "merged.xlsx")
    sheet = workbook.active
    sheet.merge_cells("A1:A2")
    sheet["A1"] = "Group A"
    sheet["B1"] = "Value"
    sheet["B2"] = "10"
    workbook.save(tmp_path / "merged.xlsx")

    sections = XlsxParser().parse(tmp_path / "merged.xlsx")
    row = next(section for section in sections if section.source_metadata["type"] == "table_row")

    assert row.source_metadata["values"] == ["Group A", "10"]
    assert row.source_metadata["merged_sources"] == {"A2": "A1"}
    assert row.source_metadata["merged_ranges"] == ["A1:A2"]


def test_formula_only_unresolved_row_is_retained_without_formula_text(tmp_path: Path) -> None:
    workbook = _save_workbook(tmp_path / "formula.xlsx")
    sheet = workbook.active
    sheet.append(["Name", "Computed"])
    sheet.append(["alpha", "=1+1"])
    sheet.append(["=2+2", None])
    workbook.save(tmp_path / "formula.xlsx")

    sections = XlsxParser().parse(tmp_path / "formula.xlsx")
    rows = [section for section in sections if section.source_metadata["type"] == "table_row"]
    unresolved = next(row for row in rows if row.source_metadata["sheet_row"] == 2)

    assert unresolved.source_metadata["formula_unresolved"] is True
    assert unresolved.source_metadata["unresolved_formula_count"] == 1
    assert unresolved.source_metadata["formula_addresses"] == ["B2"]
    assert "=1+1" not in unresolved.text
    formula_only = next(row for row in rows if row.source_metadata["sheet_row"] == 3)
    assert formula_only.source_metadata["formula_unresolved"] is True
    assert formula_only.source_metadata["formula_addresses"] == ["A3"]
    assert formula_only.source_metadata["values"] == ["", ""]


def test_cached_formula_value_wins_over_formula_expression(tmp_path: Path) -> None:
    workbook = _save_workbook(tmp_path / "cached-formula.xlsx")
    sheet = workbook.active
    sheet.append(["Name", "Computed"])
    sheet.append(["alpha", "=1+1"])
    workbook.save(tmp_path / "cached-formula.xlsx")

    source = tmp_path / "cached-formula.xlsx"
    rewritten = tmp_path / "cached-formula-rewritten.xlsx"
    with ZipFile(source) as archive, ZipFile(rewritten, "w", ZIP_DEFLATED) as output:
        for item in archive.infolist():
            content = archive.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"<f>1+1</f><v></v>", b"<f>1+1</f><v>2</v>")
            output.writestr(item, content)

    sections = XlsxParser().parse(rewritten)
    row = next(section for section in sections if section.source_metadata["type"] == "table_row")

    assert "Computed: 2" in row.text
    assert row.source_metadata["formula_unresolved"] is False
    assert row.source_metadata["formula_count"] == 1


def test_headerless_numeric_region_keeps_first_row_as_data(tmp_path: Path) -> None:
    workbook = _save_workbook(tmp_path / "headerless.xlsx")
    sheet = workbook.active
    sheet.append([1, 2])
    sheet.append([3, 4])
    workbook.save(tmp_path / "headerless.xlsx")

    sections = XlsxParser().parse(tmp_path / "headerless.xlsx")
    rows = [section for section in sections if section.source_metadata["type"] == "table_row"]

    assert not any(section.source_metadata["type"] == "table_header" for section in sections)
    assert [row.source_metadata["data_row"] for row in rows] == [1, 2]
    assert rows[0].source_metadata["headers"] == []
    assert "Columns:" not in build_search_text(
        document_name="numbers.xlsx",
        payload_text=rows[0].text,
        source_metadata=rows[0].source_metadata,
    )


def test_negative_date_and_scientific_data_rows_do_not_become_headers(tmp_path: Path) -> None:
    workbook = _save_workbook(tmp_path / "typed-data.xlsx")
    sheet = workbook.active
    sheet.append([-1, -2])
    sheet.append([-3, -4])
    sheet.append([])
    sheet.append([date(2025, 1, 1), 1e6])
    sheet.append([date(2025, 1, 2), 2e6])
    workbook.save(tmp_path / "typed-data.xlsx")

    sections = XlsxParser().parse(tmp_path / "typed-data.xlsx")
    tables = [section for section in sections if section.source_metadata["type"] == "table"]
    rows = [section for section in sections if section.source_metadata["type"] == "table_row"]

    assert len(tables) == 2
    assert all(table.source_metadata["headers"] == [] for table in tables)
    assert all(table.source_metadata["header_source"] == "none" for table in tables)
    assert [row.source_metadata["data_row"] for row in rows] == [1, 2, 1, 2]
    assert all(
        "Columns:" not in build_search_text(
            document_name="typed-data.xlsx",
            payload_text=row.text,
            source_metadata=row.source_metadata,
        )
        for row in rows
    )


def test_identifier_ip_and_time_rows_do_not_become_headers(tmp_path: Path) -> None:
    workbook = _save_workbook(tmp_path / "identifier-data.xlsx")
    sheet = workbook.active
    sheet.append(["web-01", "10.0.0.1"])
    sheet.append(["web-02", "10.0.0.2"])
    sheet.append(["web-03", "10.0.0.3"])
    sheet.append([])
    sheet.append(["09:00:00", "10:00:00"])
    sheet.append(["09:30:00", "10:30:00"])
    sheet.append(["10:00:00", "11:00:00"])
    workbook.save(tmp_path / "identifier-data.xlsx")

    sections = XlsxParser().parse(tmp_path / "identifier-data.xlsx")
    tables = [section for section in sections if section.source_metadata["type"] == "table"]
    rows = [section for section in sections if section.source_metadata["type"] == "table_row"]

    assert len(tables) == 2
    assert all(table.source_metadata["headers"] == [] for table in tables)
    assert not any(section.source_metadata["type"] == "table_header" for section in sections)
    assert [row.source_metadata["data_row"] for row in rows] == [1, 2, 3, 1, 2, 3]
    assert rows[0].text == "web-01 | 10.0.0.1"
    assert all(
        "Columns:" not in build_search_text(
            document_name="identifier-data.xlsx",
            payload_text=row.text,
            source_metadata=row.source_metadata,
        )
        for row in rows
    )


@pytest.mark.parametrize(
    ("name", "values", "explicit", "expected_headers"),
    [
        (
            "identifier-ordinary",
            [["web-01", "active"], ["web-02", "down"], ["web-03", "maintenance"]],
            False,
            [],
        ),
        (
            "ip-ordinary",
            [["10.0.0.1", "active"], ["10.0.0.2", "down"], ["10.0.0.3", "maintenance"]],
            False,
            [],
        ),
        (
            "time-ordinary",
            [["09:00:00", "active"], ["09:30:00", "down"], ["10:00:00", "maintenance"]],
            False,
            [],
        ),
        (
            "explicit-server-ip",
            [["Server", "IP Address"], ["web-01", "10.0.0.1"], ["web-02", "10.0.0.2"]],
            True,
            ["Server", "IP Address"],
        ),
        (
            "explicit-server-state",
            [["Server", "State"], ["web-01", "active"], ["web-02", "down"]],
            True,
            ["Server", "State"],
        ),
        (
            "fallback-host-state",
            [["host", "state"], ["host-01", "active"], ["host-02", "down"]],
            False,
            ["host", "state"],
        ),
        (
            "fallback-server-ip",
            [["server", "ip"], ["server-01", "10.0.0.1"], ["server-02", "10.0.0.2"]],
            False,
            ["server", "ip"],
        ),
        (
            "fallback-user-role",
            [["user", "role"], ["user-01", "admin"], ["user-02", "viewer"]],
            False,
            ["user", "role"],
        ),
        (
            "compact-aaasiems",
            [["aaasiems01", "active"], ["aaasiems02", "down"], ["aaasiems03", "maintenance"]],
            False,
            [],
        ),
        (
            "compact-ems",
            [["ems01", "active"], ["ems02", "down"], ["ems03", "maintenance"]],
            False,
            [],
        ),
        (
            "compact-host",
            [["host01", "active"], ["host02", "down"], ["host03", "maintenance"]],
            False,
            [],
        ),
    ],
)
def test_mixed_structured_and_ordinary_columns_preserve_header_contract(
    tmp_path: Path,
    name: str,
    values: list[list[str]],
    explicit: bool,
    expected_headers: list[str],
) -> None:
    from openpyxl.worksheet.table import Table

    workbook = _save_workbook(tmp_path / f"{name}.xlsx")
    sheet = workbook.active
    for row in values:
        sheet.append(row)
    if explicit:
        sheet.add_table(Table(displayName="DataTable", ref=f"A1:B{len(values)}"))
    path = tmp_path / f"{name}.xlsx"
    workbook.save(path)

    sections = XlsxParser().parse(path)
    table = next(section for section in sections if section.source_metadata["type"] == "table")
    rows = [section for section in sections if section.source_metadata["type"] == "table_row"]

    assert table.source_metadata["headers"] == expected_headers
    expected_row_count = len(values) - 1 if expected_headers else len(values)
    assert [row.source_metadata["data_row"] for row in rows] == list(
        range(1, expected_row_count + 1)
    )
    if expected_headers:
        assert "Columns: " + " | ".join(expected_headers) in build_search_text(
            document_name=f"{name}.xlsx",
            payload_text=rows[0].text,
            source_metadata=rows[0].source_metadata,
        )
    else:
        assert not any(section.source_metadata["type"] == "table_header" for section in sections)
        assert "Columns:" not in build_search_text(
            document_name=f"{name}.xlsx",
            payload_text=rows[0].text,
            source_metadata=rows[0].source_metadata,
        )


def test_header_only_explicit_table_emits_schema_sections_without_row_ranges(tmp_path: Path) -> None:
    from openpyxl.worksheet.table import Table

    workbook = _save_workbook(tmp_path / "schema-only.xlsx")
    sheet = workbook.active
    sheet.append(["Code", "State"])
    sheet.add_table(Table(displayName="SchemaOnly", ref="A1:B1"))
    workbook.save(tmp_path / "schema-only.xlsx")

    sections = XlsxParser().parse(tmp_path / "schema-only.xlsx")
    chunks = chunk_sections(uuid.uuid4(), uuid.uuid4(), sections, chunk_size=100, chunk_overlap=20)
    table = next(section for section in sections if section.source_metadata["type"] == "table")
    header = next(section for section in sections if section.source_metadata["type"] == "table_header")
    table_chunk = next(chunk for chunk in chunks if chunk.source_metadata["type"] == "table")

    assert table.source_metadata["row_count"] == 0
    assert header.text == "Code | State"
    assert not any(section.source_metadata["type"] == "table_row" for section in sections)
    assert table_chunk.source_metadata["schema_only"] is True
    assert "data_row_start" not in table_chunk.source_metadata
    assert "data_row_end" not in table_chunk.source_metadata


def test_xlsx_table_sections_enter_table_chunking_and_search_allowlist(tmp_path: Path) -> None:
    workbook = _save_workbook(tmp_path / "chunk.xlsx")
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Name", "Value"])
    sheet.append(["alpha", "10"])
    workbook.save(tmp_path / "chunk.xlsx")

    sections = XlsxParser().parse(tmp_path / "chunk.xlsx")
    chunks = chunk_sections(uuid.uuid4(), uuid.uuid4(), sections, chunk_size=100, chunk_overlap=20)
    table_chunk = next(chunk for chunk in chunks if chunk.source_metadata["type"] == "table")
    search_text = build_search_text(
        document_name="chunk.xlsx",
        payload_text=table_chunk.text,
        source_metadata=table_chunk.source_metadata,
    )

    assert table_chunk.source_metadata["table_chunk_type"] == "table"
    assert table_chunk.source_metadata["format"] == "xlsx"
    assert "Sheet: Data" in search_text
    assert "Columns: Name | Value" in search_text
    assert "formula_addresses" not in search_text


def test_xlsx_ingestion_persists_table_provenance_and_search_context(
    sqlite_session_factory,
    tmp_path: Path,
) -> None:
    workbook = _save_workbook(tmp_path / "ingest.xlsx")
    sheet = workbook.active
    sheet.title = "Inventory"
    sheet.append(["Code", "State"])
    sheet.append(["node-17", "active"])
    hidden = workbook.create_sheet("Hidden")
    hidden.append(["not-indexed"])
    workbook.save(tmp_path / "ingest.xlsx")

    class Embeddings:
        def embed_texts(self, texts):
            return [[0.1] * 1024 for _ in texts]

    with Session(sqlite_session_factory.kw["bind"]) as db:
        project = Project(name="XLSX Ingestion")
        db.add(project)
        db.flush()
        document = Document(
            project_id=project.id,
            filename="ingest.xlsx",
            storage_path=str(tmp_path / "ingest.xlsx"),
            file_size_bytes=100,
            status=DocumentStatus.uploaded,
        )
        db.add(document)
        db.flush()
        job = IngestionJob(project_id=project.id, document_id=document.id)
        db.add(job)
        db.commit()

        ingest_document_job(db, job.id, project.id, document.id, embedding_provider=Embeddings())
        db.refresh(job)
        sections = db.query(DocumentSection).all()
        chunks = db.query(Chunk).all()

    assert job.status == IngestionJobStatus.completed
    assert any(section.source_metadata["type"] == "table_row" for section in sections)
    table_chunk = next(chunk for chunk in chunks if chunk.source_metadata["type"] == "table")
    assert table_chunk.source_metadata["format"] == "xlsx"
    assert table_chunk.source_metadata["sheet_name"] == "Inventory"
    assert table_chunk.source_metadata["table_range"] == "A1:B2"
    assert "Sheet: Inventory" in table_chunk.search_text
    assert "Hidden" not in table_chunk.search_text
